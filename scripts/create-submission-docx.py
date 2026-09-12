from __future__ import annotations

from pathlib import Path
from datetime import date
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "submission-package"
OUT_DIR.mkdir(parents=True, exist_ok=True)
DIAGRAM = OUT_DIR / "matoi-unified-architecture.png"
DOCX = OUT_DIR / "Matoi_Technical_Submission_Document.docx"

PURPLE = "7C5CFF"
GREEN = "37D7A5"
DARK = "10131A"
MID = "1A1F2B"
TEXT = "E8EAEE"
MUTED = "9AA0AE"


def load_font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            pass
    return ImageFont.load_default()


def draw_wrapped(draw, text, xy, font, fill, max_width, line_gap=5):
    x, y = xy
    words = text.split()
    line = ""
    for word in words:
        test = (line + " " + word).strip()
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            line = test
        else:
            draw.text((x, y), line, font=font, fill=fill)
            y += font.size + line_gap
            line = word
    if line:
        draw.text((x, y), line, font=font, fill=fill)
        y += font.size + line_gap
    return y


def rounded_box(draw, xy, fill, outline="#3B4057", width=2, radius=18):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def make_diagram():
    W, H = 1600, 1000
    img = Image.new("RGB", (W, H), "#080A0F")
    draw = ImageDraw.Draw(img)
    title = load_font(34, True)
    subtitle = load_font(19)
    h = load_font(21, True)
    body = load_font(17)
    small = load_font(14)

    draw.text((W // 2, 38), "Matoi — Unified Agent Intelligence Market", font=title, fill="#F3F5FA", anchor="mm")
    draw.text((W // 2, 72), "One project: Arc/Circle agent treasury + Hedera x402 paid service + HCS audit proof", font=subtitle, fill="#A9AFBF", anchor="mm")

    cols = [70, 430, 790, 1150]
    labels = [
        ("Frontend / Buyer", "Next.js control room, scenario runner, Blocky402 toggle, A2A transcript"),
        ("Backend / Orchestrator", "TypeScript mission engine: provider discovery, quote ranking, spend policy, audit events"),
        ("Payment Rails", "Circle DCW on ARC-TESTNET + Hedera x402 / Blocky402-compatible facilitator"),
        ("Proof & Safety", "HCS public proofs, redacted audit log, simulation-only trading, no mainnet"),
    ]
    for i, (x, (head, desc)) in enumerate(zip(cols, labels)):
        rounded_box(draw, (x, 130, x + 300, 470), fill="#111522")
        draw.text((x + 150, 160), head, font=h, fill="#A88BFF", anchor="mm")
        draw_wrapped(draw, desc, (x + 25, 195), body, "#E8EAEE", 250, 4)

    blocks = [
        (cols[0], 285, "UI", "React / TypeScript\nArchitecture diagram\nVideo presentation"),
        (cols[1], 285, "BuyerAgent", "Mission + budget\nProvider selection\nRisk policy"),
        (cols[2], 245, "Arc/Circle", "DCW wallets\nUSDC spend auth\nNanopayments ledger"),
        (cols[2], 355, "Hedera", "x402 paid request\nUSDC 0.0.429274\nBlocky402-compatible /settle"),
        (cols[3], 285, "Evidence", "HCS topic 0.0.10426202\nReceipts x402-*\nSmoke + tests green"),
    ]
    for x, y, btitle, text in blocks:
        rounded_box(draw, (x + 30, y, x + 270, y + 85), fill="#1A2030", outline="#3A4160", radius=12)
        draw.text((x + 150, y + 22), btitle, font=small, fill="#37D7A5", anchor="mm")
        for idx, line in enumerate(text.split("\n")):
            draw.text((x + 150, y + 42 + idx * 15), line, font=small, fill="#C7CBD6", anchor="mm")

    # arrows
    arrow = "#A88BFF"
    for x1, x2 in [(370, 430), (730, 790), (1090, 1150)]:
        y = 300
        draw.line((x1, y, x2, y), fill=arrow, width=4)
        draw.polygon([(x2, y), (x2 - 14, y - 8), (x2 - 14, y + 8)], fill=arrow)

    # central flow
    rounded_box(draw, (90, 550, 1510, 870), fill="#0E121C", outline="#3A4160", radius=20)
    draw.text((W // 2, 585), "End-to-end demo flow", font=h, fill="#A88BFF", anchor="mm")
    flow = [
        "Buyer mission",
        "Provider quotes",
        "Arc/Circle spend authorization",
        "Hedera x402 paid request",
        "Receipt verified",
        "Risk review",
        "Simulation-only decision",
        "HCS proof",
    ]
    step_w = 165
    start_x = 135
    for idx, step in enumerate(flow):
        x = start_x + idx * 175
        rounded_box(draw, (x, 635, x + step_w, 735), fill="#161B29", outline="#3A4160", radius=14)
        y = draw_wrapped(draw, step, (x + 18, 663), small, "#F0F2F8", step_w - 36, 2)
        if idx < len(flow) - 1:
            ax = x + step_w
            draw.line((ax + 8, 685, ax + 35, 685), fill="#37D7A5", width=3)
            draw.polygon([(ax + 35, 685), (ax + 25, 679), (ax + 25, 691)], fill="#37D7A5")

    footer = "Verification: lint green · build green · 91 tests passed · qa:smoke OK · real paid requests verified on localhost and public tunnel"
    draw.text((W // 2, 920), footer, font=subtitle, fill="#C7CBD6", anchor="mm")
    img.save(DIAGRAM)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for k, v in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], k, bold=True, color="1F2937")
        set_cell_text(cells[1], v)
        set_cell_shading(cells[0], "F1F3F8")
    doc.add_paragraph()
    return table


def add_req_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, htxt in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], htxt, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "3B2F70")
    for row in rows:
        cells = table.add_row().cells
        for i, txt in enumerate(row):
            set_cell_text(cells[i], txt)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_doc():
    make_diagram()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Segoe UI"
    styles["Normal"].font.size = Pt(10.2)
    styles["Heading 1"].font.name = "Segoe UI Semibold"
    styles["Heading 1"].font.size = Pt(20)
    styles["Heading 1"].font.color.rgb = RGBColor(59, 47, 112)
    styles["Heading 2"].font.name = "Segoe UI Semibold"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.color.rgb = RGBColor(55, 65, 95)

    # Title page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Matoi")
    run.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor(59, 47, 112)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Agent Intelligence Market")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(31, 41, 55)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Unified Technical Submission Document")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(85, 95, 110)

    doc.add_paragraph()
    doc.add_picture(str(DIAGRAM), width=Inches(6.9))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_kv_table(doc, [
        ("Project", "Matoi — Agent Intelligence Market"),
        ("Submission tracks", "Arc/Circle Agentic Economy + Hedera x402 / Blocky402 paid service"),
        ("Contact email", "geraldwarren77@gmail.com"),
        ("X / Twitter", "@NikiYomek — https://x.com/NikiYomek"),
        ("Discord", "yomek"),
        ("Arc House", "Niki Li (yomek)"),
        ("Document date", date.today().isoformat()),
    ])

    doc.add_page_break()

    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(
        "Matoi is a control room for autonomous economic agents. A BuyerAgent opens a mission, "
        "discovers sellable intelligence providers, ranks their quotes, authorizes bounded USDC spend through "
        "Arc/Circle testnet wallets, optionally pays a Hedera x402-gated service through a Blocky402-compatible "
        "facilitator, receives normalized intelligence, produces a simulation-only trading/risk decision, and "
        "leaves a redacted public proof on Hedera HCS."
    )
    doc.add_paragraph(
        "The project is one coherent product, not two separate bounty demos: Arc/Circle is the agent treasury and "
        "USDC spend layer; Hedera is the x402 paid-service and audit-proof layer; the frontend and backend connect "
        "them into one agent mission flow."
    )

    doc.add_heading("2. Submission fit", level=1)
    add_req_table(doc, ["Track", "Requirement", "Matoi proof"], [
        ("Arc/Circle", "Functional MVP with working frontend and backend", "Next.js control room plus API backend: /api/mission/run, /api/arc/status, /api/arc/authorize-payment"),
        ("Arc/Circle", "Architecture diagram", "docs/architecture.svg and this document's embedded architecture diagram"),
        ("Arc/Circle", "Effective Circle Developer tools/tech usage", "Circle DCW role wallets, ARC-TESTNET USDC, guarded real transfer proof, spend ledger, nanopayments"),
        ("Hedera", "Live x402-gated service on Hedera testnet", "POST /api/x402/blocky402/pay"),
        ("Hedera", "Settled through Blocky402 facilitator", "Self-hosted Blocky402-compatible /settle route consumed via BLOCKY402_URL"),
        ("Hedera", "Platform/agent consumes the service", "BuyerAgent mission flow adds BLOCKY402_* receipt messages"),
        ("Hedera", "Real paid request end-to-end", "Verified localhost and public tunnel receiptIds x402-*"),
    ])

    doc.add_heading("3. Core functions", level=1)
    add_bullets(doc, [
        "Mission runner: five judge-visible buyer tasks with one-click RUN FULL MISSION flow.",
        "Provider marketplace: provider discovery, quote ranking, reputation and task-relative summaries.",
        "Arc/Circle treasury: three role wallets, spend authorization and provider payout ledger.",
        "Nanopayments: per-call metering and NANOPAYMENT_SETTLED ledger entries.",
        "Hedera x402 service: live paid request route, signed Hedera transaction payload, HMAC receipt.",
        "HCS audit: redacted, allowlisted public proof path to topic 0.0.10426202.",
        "Safety: mainnet disabled, real trades/swaps disabled, secrets never exposed in browser/audit/HCS.",
    ])

    doc.add_heading("4. Unified architecture", level=1)
    doc.add_paragraph("The following architecture is the same system diagram used in the README and presentation deck.")
    doc.add_picture(str(DIAGRAM), width=Inches(7.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Flow: Buyer mission → provider quotes → Arc/Circle spend authorization → Hedera x402 paid request → "
        "receipt verification → risk review → simulation-only decision → HCS proof."
    )

    doc.add_heading("5. Frontend and backend implementation", level=1)
    add_req_table(doc, ["Layer", "Implementation", "Key files / endpoints"], [
        ("Frontend", "Next.js / React / TypeScript control room UI", "src/components/InkExperience.tsx, route /"),
        ("Mission backend", "TypeScript orchestrator for scenarios, providers, payments, audit events", "src/lib/mission-orchestrator.ts, POST /api/mission/run"),
        ("Arc/Circle backend", "Wallet readiness, spend authorization, settlement lifecycle, optional real transfer", "GET /api/arc/status, POST /api/arc/authorize-payment, POST /api/arc/testnet-transfer/run"),
        ("Hedera backend", "x402 paid request, Blocky402-compatible facilitator, HCS audit", "POST /api/x402/blocky402/pay, /api/x402/blocky402-facilitator/settle, /api/audit/hcs"),
        ("Testing", "Unit, integration and HTTP smoke coverage", "91 tests passed, qa:smoke OK"),
    ])

    doc.add_heading("6. Circle / Arc developer technology usage", level=1)
    add_req_table(doc, ["Circle / Arc tech", "How it is used", "Proof"], [
        ("Circle Developer Controlled Wallets", "Three agent role wallets: Trader, ArcResearch, Risk", "GET /api/arc/status"),
        ("ARC-TESTNET", "All Arc wallet actions are testnet-only", "Status route and README track section"),
        ("USDC", "Agent spend authorization and optional real transfer proof", "Circle tx a0042c1c-ae9a-5120-accc-7116dfc1ec31 COMPLETE"),
        ("Circle transaction API", "Guarded real transfer endpoint with double opt-in", "POST /api/arc/testnet-transfer/run"),
        ("Agent Stack concept", "BuyerAgent, ProviderAgent, RiskAgent, ArcTreasuryAgent cooperate in one mission", "A2A transcript in /api/mission/run"),
        ("Nanopayments", "Per-provider-call metering and settlement ledger", "src/lib/nanopayments.ts; NANOPAYMENT_SETTLED entries"),
    ])

    doc.add_heading("7. Hedera x402 / Blocky402 payment flow", level=1)
    add_numbered(doc, [
        "BuyerAgent selects provider service telegram-pulse.",
        "POST /api/x402/blocky402/pay builds Hedera payment requirements: network hedera:testnet, USDC token 0.0.429274, amount 10000 smallest units = 0.01 USDC.",
        "src/lib/hedera-payment-signer.ts signs a Hedera TransferTransaction using @hashgraph/sdk and the ECDSA payer key for account 0.0.10380366.",
        "src/lib/blocky402.ts encodes the signed transaction as an x402-style PAYMENT-SIGNATURE header.",
        "src/lib/blocky402-client.ts posts to <BLOCKY402_URL>/settle.",
        "The Blocky402-compatible facilitator route auto-opens a matching challenge when needed and returns an HMAC-signed receipt.",
        "The buyer mission consumes the receipt and appends BLOCKY402_CHALLENGE_RECEIVED, BLOCKY402_PAYMENT_SUBMITTED and BLOCKY402_RECEIPT_VERIFIED messages.",
    ])
    add_kv_table(doc, [
        ("Hedera network", "testnet"),
        ("USDC token", "0.0.429274, 6 decimals"),
        ("Payer account", "0.0.10380366"),
        ("Payer EVM address", "0x4b7c9e49609650056c64bf111131503f6cca8c05"),
        ("Token association proof", "0.0.10380366@1789189919.552662668"),
        ("Recent verified receipts", "localhost x402-d72b519cbeb82c65; public tunnel x402-44df75b702d9b507"),
    ])

    doc.add_heading("8. Verification and proof state", level=1)
    add_req_table(doc, ["Check", "Result"], [
        ("npm run lint", "exit 0"),
        ("npm run build", "exit 0"),
        ("npx vitest run", "42 files, 91 tests passed"),
        ("npm run qa:smoke", "RESULT: SMOKE_HTTP_OK http://127.0.0.1:3100"),
        ("Blocky402 paid request localhost", "ok=true, receiptId=x402-d72b519cbeb82c65"),
        ("Blocky402 paid request public tunnel", "ok=true, receiptId=x402-44df75b702d9b507"),
        ("HCS topic", "0.0.10426202"),
    ])

    doc.add_heading("9. Setup and operations", level=1)
    doc.add_paragraph("Full setup lives in docs/SETUP.md. The core local run path is:")
    add_numbered(doc, [
        "npm install",
        "Copy .env.example to .env.local and fill only the integrations being tested.",
        "Start production server with BLOCKY402_URL=http://127.0.0.1:3100/api/x402/blocky402-facilitator npm run start -- --hostname 127.0.0.1 --port 3100.",
        "Expose public URL with cloudflared tunnel --url http://127.0.0.1:3100.",
        "Run npm run lint, npm run build, npx vitest run and npm run qa:smoke before submission.",
    ])

    doc.add_heading("10. Safety, privacy and non-goals", level=1)
    add_bullets(doc, [
        "Mainnet disabled; every live financial proof is testnet-only.",
        "Real trading/swaps disabled; trading desk emits simulation-only decisions.",
        "Arc real transfer path is off by default and requires ARC_REAL_USDC_TRANSFER=true plus confirmRealTransfer=true.",
        "No private keys, Circle entity secret, OpenRouter key, Telegram token, sessions or tdata are committed.",
        "HCS payloads are allowlisted and redacted; no raw private messages or credentials are published.",
        "The public demo favors honest readiness/status over hidden mocks.",
    ])

    doc.add_heading("11. Demo / presentation structure", level=1)
    add_numbered(doc, [
        "Open README and show the unified architecture diagram.",
        "Open the live control room and run a buyer mission.",
        "Show Arc/Circle role wallets, spend authorization and nanopayments.",
        "Run the Hedera x402 paid request through the public tunnel and show ok=true + receiptId.",
        "Show BLOCKY402_RECEIPT_VERIFIED in the A2A transcript.",
        "Show HCS topic 0.0.10426202 / proof path.",
        "Close with safety: testnet-only, no real trades, secrets server-side only.",
    ])

    doc.add_heading("12. Contact", level=1)
    add_kv_table(doc, [
        ("Email", "geraldwarren77@gmail.com"),
        ("X / Twitter", "@NikiYomek — https://x.com/NikiYomek"),
        ("Discord", "yomek"),
        ("Arc House", "Niki Li (yomek)"),
    ])

    doc.add_paragraph("Appendix: repository documentation")
    add_bullets(doc, [
        "README.md — judge-facing project overview and track mapping.",
        "docs/architecture.svg and docs/architecture.md — diagram and architecture details.",
        "docs/ARC_TRACK.md — Arc/Circle proof mapping.",
        "docs/HEDERA_TRACK.md — Hedera x402 / Blocky402 proof mapping.",
        "docs/SETUP.md — per-integration wiring guide.",
        "DEMO_SCRIPT.md — video demonstration shot list.",
    ])

    doc.save(DOCX)


if __name__ == "__main__":
    build_doc()
    print(DOCX)
    print(DIAGRAM)
